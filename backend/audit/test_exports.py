"""Export Excel des listes : exactement la liste filtrée, dans le périmètre, noté dans le journal."""

from datetime import date
from io import BytesIO

from django.core.cache import cache
from django.test import TestCase
from openpyxl import load_workbook

from accounts.models import User
from applications.models import Application
from applications.tests.helpers import make_application
from members.models import Member
from territories.models import Territory

from .models import AuditEntry

PASSWORD = 'un-mot-de-passe-long'
XLSX = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
JSON = 'application/json'


def territory(name, kind=Territory.Kind.COMMUNE):
    return Territory.objects.get(name=name, kind=kind)


def sheet_of(response):
    workbook = load_workbook(BytesIO(response.content))
    return workbook.active


def table(response):
    """Lignes du fichier, sans l'en-tête, et l'en-tête."""
    rows = [[cell.value for cell in row] for row in sheet_of(response).iter_rows()]
    return rows[0], rows[1:]


class ExportTestCase(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.dakar = territory('Dakar', Territory.Kind.REGION)
        cls.thies = territory('Thiès', Territory.Kind.REGION)
        cls.pikine, cls.rufisque, cls.mbour = territory('Pikine'), territory('Rufisque'), territory('Mbour')

        def person(email, first, last, role, place, joined):
            user = User.objects.create_user(email, PASSWORD, first_name=first, last_name=last, role=role, territory=place)
            return user, Member.objects.create(user=user, joined_on=joined, phone='+221 77 000 00 00', address='1, rue du Marché')

        cls.founder, _ = person('fondateur@example.org', 'Ibrahima', 'Sow', 'fondateur', None, date(2025, 12, 1))
        cls.coordinator, _ = person('coord@example.org', 'Mamadou', 'Diop', 'coordonnateur', cls.dakar, date(2026, 1, 6))
        cls.thies_coordinator, _ = person('thies@example.org', 'Khady', 'Ba', 'coordonnateur', cls.thies, date(2026, 3, 9))
        cls.awa, cls.awa_member = person('awa@example.org', 'Awa', 'Ndiaye', 'membre', cls.pikine, date(2026, 1, 15))
        cls.bineta, _ = person('bineta@example.org', 'Bineta', 'Fall', 'membre', cls.rufisque, date(2026, 2, 20))
        cls.cheikh, _ = person('cheikh@example.org', 'Cheikh', 'Sarr', 'membre', cls.pikine, date(2026, 3, 31))
        cls.dame, _ = person('dame@example.org', 'Dame', 'Thiam', 'membre', cls.mbour, date(2026, 2, 25))

    def setUp(self):
        cache.clear()


class MemberExportTests(ExportTestCase):
    url = '/api/members/export/'

    def names(self, response):
        self.assertEqual(response.status_code, 200, response.content)
        _, rows = table(response)
        return [row[2] for row in rows]  # prénoms, dans l'ordre du fichier

    def list_names(self, **params):
        response = self.client.get('/api/members/', {**params, 'page_size': 100})
        return [row['first_name'] for row in response.json()['results']]

    def test_le_fichier(self):
        self.client.force_login(self.founder)
        response = self.client.get(self.url)
        self.assertEqual(response['Content-Type'], XLSX)
        self.assertRegex(response['Content-Disposition'], r'^attachment; filename="membres-\d{4}-\d{2}-\d{2}\.xlsx"$')
        self.assertEqual(response['Cache-Control'], 'no-store')
        headers, rows = table(response)
        self.assertEqual(
            headers, ["N° d'adhérent", 'Nom', 'Prénom', 'Courriel', 'Fonction', 'Localité', "Date d'adhésion", 'Statut']
        )
        self.assertEqual(len(rows), 4)
        awa = next(row for row in rows if row[2] == 'Awa')
        self.assertEqual(awa[0], self.awa_member.number)
        self.assertEqual(awa[1:6], ['Ndiaye', 'Awa', 'awa@example.org', 'Membre simple', 'Pikine'])
        self.assertEqual(awa[6].date(), date(2026, 1, 15))
        self.assertEqual(awa[7], 'Actif')

    def test_ni_adresse_ni_telephone(self):
        self.client.force_login(self.founder)
        content = self.client.get(self.url).content
        text = ' '.join(str(cell.value) for row in sheet_of(self.client.get(self.url)).iter_rows() for cell in row)
        self.assertNotIn('rue du Marché', text)
        self.assertNotIn('+221 77', text)
        self.assertTrue(content)

    def test_le_fichier_est_exactement_la_liste_filtree(self):
        """Avec n'importe quels filtres, le fichier contient les mêmes personnes, dans le même ordre, que la liste."""
        self.client.force_login(self.founder)
        cases = [
            {},
            {'territory': self.dakar.pk},
            {'territory': self.pikine.pk, 'joined_from': '2026-01-01', 'joined_to': '2026-03-30'},
            {'joined_from': '2026-02-01', 'joined_to': '2026-02-28'},
            {'search': 'ndiaye'},
            {'ordering': 'name'},
            {'ordering': '-place'},
            {'territory': self.thies.pk, 'status': 'actif'},
            {'role': 'coordonnateur,fondateur'},
            {'role': 'coordonnateur,fondateur', 'territory': self.dakar.pk, 'ordering': 'date'},
        ]
        for params in cases:
            with self.subTest(params=params):
                self.assertEqual(self.names(self.client.get(self.url, params)), self.list_names(**params))

    def test_filtres_sans_resultat_donne_un_fichier_avec_les_seuls_en_tetes(self):
        self.client.force_login(self.founder)
        headers, rows = table(self.client.get(self.url, {'search': 'inconnu'}))
        self.assertEqual(len(headers), 8)
        self.assertEqual(rows, [])

    def test_un_coordonnateur_n_exporte_que_son_perimetre(self):
        self.client.force_login(self.coordinator)  # région de Dakar
        self.assertEqual(sorted(self.names(self.client.get(self.url))), ['Awa', 'Bineta', 'Cheikh'])
        self.assertEqual(self.names(self.client.get(self.url, {'territory': self.thies.pk})), [])
        self.assertEqual(self.names(self.client.get(self.url, {'territory': self.mbour.pk})), [])

    def test_base_des_coordonnateurs(self):
        self.client.force_login(self.founder)
        response = self.client.get(self.url, {'role': 'coordonnateur,fondateur'})
        self.assertIn('coordonnateurs-', response['Content-Disposition'])
        headers, rows = table(response)
        self.assertIn("Date d'entrée", headers)
        self.assertEqual({row[4] for row in rows}, {'Coordonnateur', 'Membre fondateur'})
        self.assertIn('Tous les territoires', {row[5] for row in rows})

    def test_les_suspendus_figurent_avec_leur_statut(self):
        self.client.force_login(self.founder)
        self.client.post(f'/api/members/{self.awa_member.pk}/status/', {'status': 'suspendu', 'reason': 'Motif.'}, content_type=JSON)
        _, rows = table(self.client.get(self.url, {'status': 'suspendu'}))
        self.assertEqual([(row[2], row[7]) for row in rows], [('Awa', 'Suspendu')])

    def test_une_formule_saisie_dans_un_nom_reste_du_texte(self):
        self.awa.first_name = '=HYPERLINK("http://example.org","Cliquez")'
        self.awa.last_name = '+SUM(1+1)'
        self.awa.save()
        self.client.force_login(self.founder)
        sheet = sheet_of(self.client.get(self.url, {'search': 'SUM'}))
        first, last = sheet['C2'], sheet['B2']
        self.assertEqual(first.value, '=HYPERLINK("http://example.org","Cliquez")')
        self.assertEqual((first.data_type, last.data_type), ('s', 's'))
        import zipfile

        archive = zipfile.ZipFile(BytesIO(self.client.get(self.url, {'search': 'SUM'}).content))
        self.assertNotIn('<f>', archive.read('xl/worksheets/sheet1.xml').decode())

    def test_reserve_au_bureau(self):
        self.client.force_login(self.awa)
        self.assertEqual(self.client.get(self.url).status_code, 403)
        self.client.logout()
        self.assertEqual(self.client.get(self.url).status_code, 401)
        self.assertFalse(AuditEntry.objects.exists())

    def test_filtre_invalide_ni_fichier_ni_journal(self):
        self.client.force_login(self.founder)
        for params in ({'territory': 'abc'}, {'joined_from': 'hier'}, {'status': 'banni'}, {'role': 'roi'}, {'ordering': 'age'}):
            self.assertEqual(self.client.get(self.url, params).status_code, 400, params)
        self.assertFalse(AuditEntry.objects.exists())

    def test_methodes_d_ecriture_refusees(self):
        self.client.force_login(self.founder)
        self.assertEqual(self.client.post(self.url, {}, content_type=JSON).status_code, 405)


class ExportJournalTests(ExportTestCase):
    def test_chaque_export_est_note_dans_le_journal(self):
        self.client.force_login(self.coordinator)
        response = self.client.get(
            '/api/members/export/', {'territory': self.pikine.pk, 'status': 'actif', 'joined_from': '2026-01-01', 'search': 'ndiaye'}
        )
        self.assertEqual(response.status_code, 200)
        entry = AuditEntry.objects.get()
        self.assertEqual(entry.action, 'liste_exportee')
        self.assertEqual((entry.actor, entry.actor_name, entry.subject), (self.coordinator, 'Mamadou Diop', 'Base des membres'))
        self.assertEqual(entry.details['rows'], 1)
        self.assertEqual(
            entry.details['filters'],
            {'Fonction': 'Membre simple', 'Statut': 'Actif', 'Localité': 'Pikine', 'Du': '2026-01-01', 'Recherche': 'ndiaye'},
        )
        self.assertRegex(entry.details['file'], r'^membres-\d{4}-\d{2}-\d{2}\.xlsx$')

    def test_chaque_export_a_sa_ligne(self):
        self.client.force_login(self.founder)
        for _ in range(3):
            self.client.get('/api/members/export/')
        self.assertEqual(AuditEntry.objects.filter(action='liste_exportee').count(), 3)

    def test_le_journal_montre_l_export_aux_fondateurs(self):
        self.client.force_login(self.coordinator)
        self.client.get('/api/members/export/', {'territory': self.pikine.pk})
        self.client.force_login(self.founder)
        row = self.client.get('/api/audit/').json()['results'][0]
        self.assertEqual((row['action'], row['subject']), ('liste_exportee', 'Base des membres'))
        self.assertEqual(row['detail'], '2 lignes · Fonction : Membre simple, Localité : Pikine')
        self.assertEqual(row['actor']['name'], 'Mamadou Diop')

    def test_l_export_des_coordonnateurs_est_note_sous_son_nom(self):
        self.client.force_login(self.founder)
        self.client.get('/api/members/export/', {'role': 'coordonnateur,fondateur'})
        entry = AuditEntry.objects.get()
        self.assertEqual(entry.subject, 'Base des coordonnateurs')
        self.assertEqual(entry.details['filters']['Fonction'], 'Coordonnateur, Membre fondateur')
        self.assertEqual(entry.details['rows'], 3)


class ApplicationExportTests(ExportTestCase):
    url = '/api/applications/export/'

    @classmethod
    def setUpTestData(cls):
        super().setUpTestData()
        cls.a = make_application(territory('Podor'), email='a@example.org', first_name='Awa', last_name='Sy', motivation='Texte secret de motivation')
        cls.b = make_application(territory('Pikine'), email='b@example.org', first_name='Omar', last_name='Niang', status=Application.Status.REJETEE, rejection_reason='Motif secret')

    def test_le_fichier_est_la_liste_filtree(self):
        self.client.force_login(self.founder)
        for params in ({}, {'status': 'rejetee'}, {'search': 'sy'}, {'status': 'a_examiner', 'search': 'omar'}):
            with self.subTest(params=params):
                response = self.client.get(self.url, params)
                self.assertEqual(response.status_code, 200)
                _, rows = table(response)
                listed = self.client.get('/api/applications/', params).json()['results']
                self.assertEqual([row[0] for row in rows], [item['reference'] for item in listed])

    def test_colonnes_sans_donnees_sensibles(self):
        self.client.force_login(self.founder)
        response = self.client.get(self.url)
        headers, rows = table(response)
        self.assertEqual(
            headers,
            ['Référence', 'Nom', 'Prénom', 'Courriel', 'Localité souhaitée', 'Profession', 'Statut', 'Date de la demande', 'Date de la décision', 'Décidée par'],
        )
        text = ' '.join(str(cell) for row in rows for cell in row)
        self.assertNotIn('secret', text)
        self.assertIn('Rejetée', text)

    def test_note_dans_le_journal(self):
        self.client.force_login(self.founder)
        self.client.get(self.url, {'status': 'rejetee'})
        entry = AuditEntry.objects.get()
        self.assertEqual((entry.action, entry.subject), ('liste_exportee', "Demandes d'adhésion"))
        self.assertEqual(entry.details['rows'], 1)
        self.assertEqual(entry.details['filters'], {'Statut': 'Rejetée'})

    def test_reserve_aux_fondateurs(self):
        self.client.force_login(self.coordinator)
        self.assertEqual(self.client.get(self.url).status_code, 403)
        self.client.logout()
        self.assertEqual(self.client.get(self.url).status_code, 401)
        self.assertFalse(AuditEntry.objects.exists())

    def test_filtre_inconnu(self):
        self.client.force_login(self.founder)
        self.assertEqual(self.client.get(self.url, {'status': 'banni'}).status_code, 400)
        self.assertFalse(AuditEntry.objects.exists())


class JournalExportTests(ExportTestCase):
    url = '/api/audit/export/'

    def test_le_journal_s_exporte_avec_ses_filtres_et_l_export_est_note(self):
        self.client.force_login(self.founder)
        self.client.post(f'/api/members/{self.awa_member.pk}/status/', {'status': 'suspendu', 'reason': 'Motif.'}, content_type=JSON)
        self.client.get('/api/members/export/')
        response = self.client.get(self.url, {'action': 'statut_modifie'})
        self.assertEqual(response.status_code, 200)
        headers, rows = table(response)
        self.assertEqual(headers, ['Date', 'Action', 'Concerne', 'Détail', 'Auteur'])
        self.assertEqual([(row[1], row[2], row[4]) for row in rows], [('Statut modifié', 'Awa Ndiaye', 'Ibrahima Sow')])
        self.assertIn('Motif : Motif.', rows[0][3])
        # L'export du journal est lui aussi noté, mais il n'est pas dans le fichier qu'il vient de produire.
        entry = AuditEntry.objects.filter(subject='Journal des actions').get()
        self.assertEqual(entry.details['filters'], {'Action': 'Statut modifié'})
        self.assertEqual(entry.details['rows'], 1)

    def test_reserve_aux_fondateurs(self):
        self.client.force_login(self.coordinator)
        self.assertEqual(self.client.get(self.url).status_code, 403)
        self.assertFalse(AuditEntry.objects.exists())
