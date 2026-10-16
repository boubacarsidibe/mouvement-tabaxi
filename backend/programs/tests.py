from datetime import timedelta
from io import BytesIO

from django.core.cache import cache
from django.db.models import ProtectedError
from django.test import TestCase
from django.utils import timezone
from openpyxl import load_workbook

from accounts.models import User
from audit.models import AuditEntry
from members.models import Member
from territories.models import Territory

from .models import Activity, Attendance, Program, Registration

PASSWORD = 'un-mot-de-passe-long'
JSON = 'application/json'


def territory(name, kind=Territory.Kind.COMMUNE):
    return Territory.objects.get(name=name, kind=kind)


def days(count):
    return timezone.now() + timedelta(days=count)


class ProgramsTestCase(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.dakar = territory('Dakar', Territory.Kind.REGION)
        cls.thies = territory('Thiès', Territory.Kind.REGION)
        cls.pikine, cls.rufisque, cls.mbour = territory('Pikine'), territory('Rufisque'), territory('Mbour')

        def person(email, first, last, role, place):
            user = User.objects.create_user(email, PASSWORD, first_name=first, last_name=last, role=role, territory=place)
            return user, Member.objects.create(user=user)

        cls.founder, cls.founder_member = person('fondateur@example.org', 'Ibrahima', 'Sow', 'fondateur', None)
        cls.coord_dakar, _ = person('dakar@example.org', 'Mamadou', 'Diop', 'coordonnateur', cls.dakar)
        cls.coord_pikine, _ = person('pikine@example.org', 'Rokhaya', 'Gueye', 'coordonnateur', cls.pikine)
        cls.coord_thies, _ = person('thies@example.org', 'Khady', 'Ba', 'coordonnateur', cls.thies)
        cls.awa, cls.awa_member = person('awa@example.org', 'Awa', 'Ndiaye', 'membre', cls.pikine)
        cls.bineta, cls.bineta_member = person('bineta@example.org', 'Bineta', 'Fall', 'membre', cls.pikine)
        cls.cheikh, cls.cheikh_member = person('cheikh@example.org', 'Cheikh', 'Sarr', 'membre', cls.rufisque)
        cls.dame, cls.dame_member = person('dame@example.org', 'Dame', 'Thiam', 'membre', cls.mbour)
        cls.fatou, cls.fatou_member = person('fatou@example.org', 'Fatou', 'Diallo', 'membre', cls.pikine)
        Member.objects.filter(pk=cls.fatou_member.pk).update(status=Member.Status.SUSPENDU)

        today = timezone.localdate()
        cls.pikine_program = Program.objects.create(
            name='Brigades citoyennes', objective='Mobiliser les jeunes.', territory=cls.pikine,
            manager=cls.coord_pikine, start_date=today - timedelta(days=30), end_date=today + timedelta(days=60),
            status=Program.Status.ONGOING, created_by=cls.coord_pikine,
        )
        cls.thies_program = Program.objects.create(
            name='Éducation civique', objective='Sensibiliser dans les écoles.', territory=cls.mbour,
            manager=cls.coord_thies, start_date=today, end_date=today + timedelta(days=90), created_by=cls.coord_thies,
        )
        cls.region_program = Program.objects.create(
            name='Santé communautaire', objective='Prévention.', territory=cls.dakar,
            start_date=today - timedelta(days=10), end_date=today + timedelta(days=40), status=Program.Status.PLANNED,
        )
        cls.past_activity = Activity.objects.create(
            program=cls.pikine_program, title='Nettoyage de la plage', starts_at=days(-3), ends_at=days(-3) + timedelta(hours=3),
            place='Plage de Yoff', description='Collecte des déchets.',
        )
        cls.future_activity = Activity.objects.create(
            program=cls.pikine_program, title='Atelier de quartier', starts_at=days(5), place='Mairie', description='Échanges avec les riverains.',
        )
        cls.thies_activity = Activity.objects.create(
            program=cls.thies_program, title='Causerie à l\'école', starts_at=days(7), place='École de Mbour', description='Éducation civique.',
        )

    def setUp(self):
        cache.clear()

    def login(self, user):
        self.client.force_login(user)

    def post(self, url, payload=None):
        return self.client.post(url, payload or {}, content_type=JSON)

    def patch(self, url, payload):
        return self.client.patch(url, payload, content_type=JSON)

    def put(self, url, payload):
        return self.client.put(url, payload, content_type=JSON)

    def program_payload(self, **changes):
        today = timezone.localdate()
        return {
            'name': 'Lecture pour tous', 'objective': 'Donner le goût de lire.', 'territory': self.pikine.pk,
            'manager': self.coord_pikine.pk, 'start_date': today.isoformat(),
            'end_date': (today + timedelta(days=30)).isoformat(), 'status': 'prevu', **changes,
        }


class ProgramCreationTests(ProgramsTestCase):
    url = '/api/programs/'

    def test_un_coordonnateur_cree_un_programme_pour_son_territoire(self):
        self.login(self.coord_pikine)
        response = self.post(self.url, self.program_payload())
        self.assertEqual(response.status_code, 201, response.content)
        data = response.json()
        self.assertEqual((data['name'], data['status'], data['status_label']), ('Lecture pour tous', 'prevu', 'Prévu'))
        self.assertEqual(data['territory']['name'], 'Pikine')
        self.assertEqual(data['manager']['name'], 'Rokhaya Gueye')
        self.assertEqual(data['activities'], [])
        program = Program.objects.get(pk=data['id'])
        self.assertEqual(program.created_by, self.coord_pikine)

    def test_le_coordonnateur_d_une_region_gere_ses_communes(self):
        self.login(self.coord_dakar)
        self.assertEqual(self.post(self.url, self.program_payload(territory=self.rufisque.pk, manager=None)).status_code, 201)
        self.assertEqual(self.post(self.url, self.program_payload(territory=self.dakar.pk, manager=self.coord_dakar.pk)).status_code, 201)

    def test_un_coordonnateur_ne_cree_pas_hors_de_son_territoire(self):
        self.login(self.coord_pikine)
        for place in (self.rufisque, self.mbour, self.dakar):
            response = self.post(self.url, self.program_payload(territory=place.pk, manager=None))
            self.assertEqual(response.status_code, 400, place.name)
            self.assertIn('hors de votre périmètre', response.json()['territory'][0])
        self.assertEqual(Program.objects.count(), 3)

    def test_un_fondateur_cree_partout(self):
        self.login(self.founder)
        for place in (self.pikine, self.mbour, self.dakar):
            response = self.post(self.url, self.program_payload(territory=place.pk, manager=self.founder.pk))
            self.assertEqual(response.status_code, 201, place.name)

    def test_les_quatre_etats(self):
        self.login(self.founder)
        for status, label in (('prevu', 'Prévu'), ('en_cours', 'En cours'), ('termine', 'Terminé'), ('annule', 'Annulé')):
            response = self.post(self.url, self.program_payload(status=status, manager=None))
            self.assertEqual((response.status_code, response.json()['status_label']), (201, label))
        self.assertEqual(self.post(self.url, self.program_payload(status='archive')).status_code, 400)

    def test_etat_par_defaut_prevu(self):
        self.login(self.founder)
        payload = self.program_payload(manager=None)
        del payload['status']
        self.assertEqual(self.post(self.url, payload).json()['status'], 'prevu')

    def test_champs_obligatoires_en_francais(self):
        self.login(self.founder)
        response = self.post(self.url, {})
        self.assertEqual(response.status_code, 400)
        errors = response.json()
        self.assertEqual(errors['name'], ['Donnez un nom au programme.'])
        self.assertEqual(errors['objective'], ["Précisez l'objectif du programme."])
        self.assertEqual(errors['territory'], ['Choisissez la localité concernée.'])
        self.assertEqual(errors['start_date'], ['Indiquez la date de début.'])
        self.assertEqual(errors['end_date'], ['Indiquez la date de fin.'])

    def test_fin_avant_debut(self):
        self.login(self.founder)
        response = self.post(self.url, self.program_payload(start_date='2026-06-10', end_date='2026-06-01', manager=None))
        self.assertEqual(response.status_code, 400)
        self.assertEqual(response.json()['end_date'], ['La date de fin doit être égale ou postérieure à la date de début.'])
        self.assertEqual(self.post(self.url, self.program_payload(start_date='2026-06-10', end_date='2026-06-10', manager=None)).status_code, 201)

    def test_dates_invalides(self):
        self.login(self.founder)
        response = self.post(self.url, self.program_payload(start_date='10/06/2026', manager=None))
        self.assertEqual(response.json()['start_date'], ['Date invalide : utilisez le format AAAA-MM-JJ.'])

    def test_localite_inconnue_ou_inactive(self):
        self.login(self.founder)
        self.assertEqual(self.post(self.url, self.program_payload(territory=999999, manager=None)).json()['territory'], ['Localité inconnue.'])
        Territory.objects.filter(pk=self.rufisque.pk).update(is_active=False)
        response = self.post(self.url, self.program_payload(territory=self.rufisque.pk, manager=None))
        self.assertEqual(response.json()['territory'], ["Cette localité n'est plus active."])

    def test_le_responsable_coordonne_la_localite(self):
        self.login(self.founder)
        # Un coordonnateur d'une autre région : refusé.
        response = self.post(self.url, self.program_payload(manager=self.coord_thies.pk))
        self.assertEqual(response.status_code, 400)
        self.assertIn('ne coordonne pas cette localité', response.json()['manager'][0])
        # Le coordonnateur de la région, qui gère aussi la commune : accepté ; un fondateur aussi.
        self.assertEqual(self.post(self.url, self.program_payload(manager=self.coord_dakar.pk)).status_code, 201)
        self.assertEqual(self.post(self.url, self.program_payload(manager=self.founder.pk)).status_code, 201)
        # Un membre simple ne peut pas être responsable ; un compte désactivé non plus.
        self.assertEqual(self.post(self.url, self.program_payload(manager=self.awa.pk)).status_code, 400)
        User.objects.filter(pk=self.coord_pikine.pk).update(is_active=False)
        self.assertEqual(self.post(self.url, self.program_payload(manager=self.coord_pikine.pk)).status_code, 400)
        self.assertEqual(self.post(self.url, self.program_payload(manager=99999)).json()['manager'], ['Responsable inconnu.'])

    def test_sans_responsable(self):
        self.login(self.coord_pikine)
        response = self.post(self.url, self.program_payload(manager=None))
        self.assertEqual((response.status_code, response.json()['manager']), (201, None))

    def test_reserve_au_bureau(self):
        self.login(self.awa)
        self.assertEqual(self.post(self.url, self.program_payload()).status_code, 403)
        self.assertEqual(self.client.get(self.url).status_code, 403)
        self.client.logout()
        self.assertEqual(self.post(self.url, self.program_payload()).status_code, 401)
        self.assertEqual(Program.objects.count(), 3)


class ProgramListAndDetailTests(ProgramsTestCase):
    url = '/api/programs/'

    def names(self, **params):
        response = self.client.get(self.url, params)
        self.assertEqual(response.status_code, 200, response.content)
        return [row['name'] for row in response.json()['results']]

    def test_un_fondateur_voit_tous_les_programmes(self):
        self.login(self.founder)
        self.assertEqual(sorted(self.names()), ['Brigades citoyennes', 'Santé communautaire', 'Éducation civique'])

    def test_un_coordonnateur_ne_voit_que_son_territoire(self):
        self.login(self.coord_pikine)
        self.assertEqual(self.names(), ['Brigades citoyennes'])  # pas le programme de la région, ni celui de Mbour
        self.login(self.coord_dakar)
        self.assertEqual(sorted(self.names()), ['Brigades citoyennes', 'Santé communautaire'])
        self.login(self.coord_thies)
        self.assertEqual(self.names(), ['Éducation civique'])

    def test_filtres(self):
        self.login(self.founder)
        self.assertEqual(self.names(status='en_cours'), ['Brigades citoyennes'])
        self.assertEqual(self.names(status='annule'), [])
        self.assertEqual(sorted(self.names(territory=self.dakar.pk)), ['Brigades citoyennes', 'Santé communautaire'])
        self.assertEqual(self.names(territory=self.mbour.pk), ['Éducation civique'])
        self.assertEqual(self.names(search='écoles'), ['Éducation civique'])
        self.assertEqual(self.names(search='brigades jeunes'), ['Brigades citoyennes'])
        self.assertEqual(self.names(search='brigades écoles'), [])
        for params in ({'status': 'archive'}, {'territory': 'abc'}, {'ordering': 'age'}):
            self.assertEqual(self.client.get(self.url, params).status_code, 400, params)

    def test_un_coordonnateur_ne_sort_pas_de_son_perimetre_avec_un_filtre(self):
        self.login(self.coord_pikine)
        self.assertEqual(self.names(territory=self.mbour.pk), [])

    def test_tri_et_pages(self):
        self.login(self.founder)
        self.assertEqual(self.names(ordering='name'), ['Brigades citoyennes', 'Santé communautaire', 'Éducation civique'])
        self.assertEqual(self.names(ordering='-name')[0], 'Éducation civique')
        data = self.client.get(self.url, {'page_size': 2}).json()
        self.assertEqual((data['count'], len(data['results'])), (3, 2))

    def test_la_ligne_du_programme(self):
        self.login(self.coord_pikine)
        row = self.client.get(self.url).json()['results'][0]
        self.assertEqual(row['activities_count'], 2)
        self.assertEqual((row['status'], row['manager']['name'], row['territory']['name']), ('en_cours', 'Rokhaya Gueye', 'Pikine'))
        self.assertNotIn('activities', row)

    def test_la_fiche_liste_les_activites_par_ordre_de_date(self):
        later = Activity.objects.create(
            program=self.pikine_program, title='Bilan', starts_at=days(20), place='Mairie', description='Bilan du programme.'
        )
        Activity.objects.create(
            program=self.pikine_program, title='Lancement', starts_at=days(-20), place='Mairie', description='Lancement.'
        )
        self.login(self.coord_pikine)
        data = self.client.get(f'{self.url}{self.pikine_program.pk}/').json()
        self.assertEqual([a['title'] for a in data['activities']], ['Lancement', 'Nettoyage de la plage', 'Atelier de quartier', 'Bilan'])
        self.assertEqual(data['activities_count'], 4)
        self.assertEqual(data['activities'][-1]['id'], later.pk)

    def test_la_fiche_hors_perimetre_n_existe_pas(self):
        self.login(self.coord_pikine)
        for program in (self.thies_program, self.region_program):
            self.assertEqual(self.client.get(f'{self.url}{program.pk}/').status_code, 404)
            self.assertEqual(self.patch(f'{self.url}{program.pk}/', {'name': 'Piraté'}).status_code, 404)
        self.assertEqual(Program.objects.get(pk=self.thies_program.pk).name, 'Éducation civique')

    def test_pas_de_suppression(self):
        self.login(self.founder)
        self.assertEqual(self.client.delete(f'{self.url}{self.pikine_program.pk}/').status_code, 405)
        self.assertTrue(Program.objects.filter(pk=self.pikine_program.pk).exists())

    def test_un_territoire_qui_a_un_programme_ne_se_supprime_pas(self):
        with self.assertRaises(ProtectedError):
            self.pikine.delete()


class ProgramUpdateTests(ProgramsTestCase):
    def url(self, program):
        return f'/api/programs/{program.pk}/'

    def test_un_coordonnateur_modifie_son_programme(self):
        self.login(self.coord_pikine)
        response = self.patch(self.url(self.pikine_program), {'name': 'Brigades de quartier', 'status': 'termine'})
        self.assertEqual(response.status_code, 200, response.content)
        self.assertEqual((response.json()['name'], response.json()['status_label']), ('Brigades de quartier', 'Terminé'))
        self.assertEqual(len(response.json()['activities']), 2)

    def test_un_fondateur_modifie_tous_les_programmes(self):
        self.login(self.founder)
        for program in (self.pikine_program, self.thies_program, self.region_program):
            self.assertEqual(self.patch(self.url(program), {'status': 'annule'}).status_code, 200)

    def test_on_ne_deplace_pas_un_programme_hors_de_son_perimetre(self):
        self.login(self.coord_pikine)
        response = self.patch(self.url(self.pikine_program), {'territory': self.rufisque.pk, 'manager': None})
        self.assertEqual(response.status_code, 400)
        self.assertEqual(Program.objects.get(pk=self.pikine_program.pk).territory, self.pikine)

    def test_le_responsable_doit_toujours_coordonner_la_localite(self):
        self.login(self.founder)
        response = self.patch(self.url(self.pikine_program), {'territory': self.mbour.pk})
        self.assertEqual(response.status_code, 400)  # Rokhaya (Pikine) ne coordonne pas Mbour
        self.assertEqual(self.patch(self.url(self.pikine_program), {'territory': self.mbour.pk, 'manager': self.coord_thies.pk}).status_code, 200)

    def test_dates_incoherentes(self):
        self.login(self.coord_pikine)
        response = self.patch(self.url(self.pikine_program), {'end_date': (self.pikine_program.start_date - timedelta(days=1)).isoformat()})
        self.assertEqual(response.status_code, 400)

    def test_on_ne_resserre_pas_les_dates_autour_d_activites_qui_sortent(self):
        self.login(self.coord_pikine)
        # L'atelier a lieu dans 5 jours : une fin de programme demain le laisserait dehors.
        tomorrow = (timezone.localdate() + timedelta(days=1)).isoformat()
        response = self.patch(self.url(self.pikine_program), {'end_date': tomorrow})
        self.assertEqual(response.status_code, 400)
        self.assertIn('activité(s) se déroulent en dehors de ces dates', response.json()['start_date'][0])
        # Élargir est toujours possible.
        later = (timezone.localdate() + timedelta(days=200)).isoformat()
        self.assertEqual(self.patch(self.url(self.pikine_program), {'end_date': later}).status_code, 200)

    def test_un_membre_ne_modifie_rien(self):
        self.login(self.awa)
        self.assertEqual(self.patch(self.url(self.pikine_program), {'name': 'x'}).status_code, 403)

    def test_managers(self):
        self.login(self.coord_pikine)
        response = self.client.get('/api/programs/managers/', {'territory': self.pikine.pk})
        self.assertEqual(response.status_code, 200)
        names = [m['name'] for m in response.json()]
        self.assertEqual(names, ['Mamadou Diop', 'Rokhaya Gueye', 'Ibrahima Sow'])  # tri par nom de famille
        self.assertEqual(self.client.get('/api/programs/managers/', {'territory': self.mbour.pk}).status_code, 404)
        self.assertEqual(self.client.get('/api/programs/managers/', {'territory': 'x'}).status_code, 400)
        self.login(self.awa)
        self.assertEqual(self.client.get('/api/programs/managers/', {'territory': self.pikine.pk}).status_code, 403)


class ActivityTests(ProgramsTestCase):
    def url(self, program):
        return f'/api/programs/{program.pk}/activities/'

    def payload(self, **changes):
        start = timezone.localtime(days(10)).replace(hour=9, minute=0, second=0, microsecond=0)
        return {
            'title': 'Distribution de livres', 'starts_at': start.isoformat(),
            'ends_at': (start + timedelta(hours=3)).isoformat(), 'place': 'Bibliothèque de Pikine',
            'description': 'Distribution de livres aux enfants.', **changes,
        }

    def test_ajout_d_une_activite(self):
        self.login(self.coord_pikine)
        response = self.post(self.url(self.pikine_program), self.payload())
        self.assertEqual(response.status_code, 201, response.content)
        data = response.json()
        self.assertEqual((data['title'], data['place']), ('Distribution de livres', 'Bibliothèque de Pikine'))
        self.assertEqual(data['program']['name'], 'Brigades citoyennes')
        self.assertEqual((data['registrations_count'], data['attendance_noted'], data['attendance_count']), (0, False, None))
        self.assertEqual(Activity.objects.get(pk=data['id']).created_by, self.coord_pikine)

    def test_sans_heure_de_fin(self):
        self.login(self.coord_pikine)
        payload = self.payload()
        del payload['ends_at']
        self.assertEqual(self.post(self.url(self.pikine_program), payload).status_code, 201)

    def test_champs_obligatoires_en_francais(self):
        self.login(self.coord_pikine)
        errors = self.post(self.url(self.pikine_program), {}).json()
        self.assertEqual(errors['title'], ["Donnez un titre à l'activité."])
        self.assertEqual(errors['place'], ["Indiquez le lieu de l'activité."])
        self.assertEqual(errors['description'], ["Décrivez l'activité."])
        self.assertEqual(errors['starts_at'], ["Indiquez la date et l'heure de l'activité."])

    def test_l_activite_se_deroule_pendant_le_programme(self):
        self.login(self.coord_pikine)
        outside = days(-60).isoformat()
        response = self.post(self.url(self.pikine_program), self.payload(starts_at=outside, ends_at=None))
        self.assertEqual(response.status_code, 400)
        self.assertIn('doit se situer entre le', response.json()['starts_at'][0])
        response = self.post(self.url(self.pikine_program), self.payload(ends_at=days(400).isoformat()))
        self.assertEqual(response.status_code, 400)
        self.assertIn('doit finir au plus tard le', response.json()['ends_at'][0])

    def test_fin_avant_debut(self):
        self.login(self.coord_pikine)
        start = days(10)
        response = self.post(self.url(self.pikine_program), self.payload(starts_at=start.isoformat(), ends_at=(start - timedelta(hours=1)).isoformat()))
        self.assertEqual(response.json()['ends_at'], ["L'heure de fin doit suivre l'heure de début."])

    def test_pas_d_activite_dans_un_programme_annule(self):
        Program.objects.filter(pk=self.pikine_program.pk).update(status=Program.Status.CANCELLED)
        self.login(self.coord_pikine)
        response = self.post(self.url(self.pikine_program), self.payload())
        self.assertEqual(response.status_code, 400)
        self.assertIn('annulé', response.json()['detail'][0])

    def test_hors_perimetre(self):
        self.login(self.coord_pikine)
        self.assertEqual(self.post(self.url(self.thies_program), self.payload()).status_code, 404)
        self.assertEqual(self.client.get(f'/api/activities/{self.thies_activity.pk}/').status_code, 404)
        self.assertEqual(self.patch(f'/api/activities/{self.thies_activity.pk}/', {'title': 'x'}).status_code, 404)
        self.login(self.founder)
        self.assertEqual(self.post(self.url(self.thies_program), self.payload()).status_code, 201)

    def test_modification(self):
        self.login(self.coord_pikine)
        response = self.patch(f'/api/activities/{self.future_activity.pk}/', {'place': 'Salle des fêtes'})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()['place'], 'Salle des fêtes')
        response = self.patch(f'/api/activities/{self.future_activity.pk}/', {'title': ''})
        self.assertEqual(response.status_code, 400)

    def test_un_membre_ne_planifie_rien(self):
        self.login(self.awa)
        self.assertEqual(self.post(self.url(self.pikine_program), self.payload()).status_code, 403)
        self.assertEqual(self.patch(f'/api/activities/{self.future_activity.pk}/', {'title': 'x'}).status_code, 403)


class ReportTests(ProgramsTestCase):
    def url(self, activity):
        return f'/api/activities/{activity.pk}/report/'

    def test_compte_rendu_apres_l_activite(self):
        self.login(self.coord_pikine)
        response = self.put(self.url(self.past_activity), {'report': '  40 participants, 12 sacs ramassés.  '})
        self.assertEqual(response.status_code, 200, response.content)
        data = response.json()
        self.assertEqual(data['report'], '40 participants, 12 sacs ramassés.')
        self.assertEqual(data['report_author']['name'], 'Rokhaya Gueye')
        self.assertIsNotNone(data['reported_at'])

    def test_le_compte_rendu_est_rattache_a_l_activite_et_visible_du_bureau_du_territoire(self):
        self.login(self.coord_pikine)
        self.put(self.url(self.past_activity), {'report': 'Bilan de la journée.'})
        for viewer in (self.coord_dakar, self.founder, self.coord_pikine):
            self.login(viewer)
            self.assertEqual(self.client.get(f'/api/activities/{self.past_activity.pk}/').json()['report'], 'Bilan de la journée.')
            program = self.client.get(f'/api/programs/{self.pikine_program.pk}/').json()
            self.assertIn('Bilan de la journée.', [a['report'] for a in program['activities']])
        self.login(self.coord_thies)  # un autre territoire ne le voit pas
        self.assertEqual(self.client.get(f'/api/activities/{self.past_activity.pk}/').status_code, 404)

    def test_pas_de_compte_rendu_avant_l_activite(self):
        self.login(self.coord_pikine)
        response = self.put(self.url(self.future_activity), {'report': 'Trop tôt.'})
        self.assertEqual(response.status_code, 409)
        self.assertEqual(response.json()['code'], 'not_started')
        self.assertEqual(Activity.objects.get(pk=self.future_activity.pk).report, '')

    def test_met_a_jour_l_etat_du_programme(self):
        self.login(self.coord_pikine)
        response = self.put(self.url(self.past_activity), {'report': 'Programme terminé avec succès.', 'program_status': 'termine'})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()['program']['status_label'], 'Terminé')
        self.assertEqual(Program.objects.get(pk=self.pikine_program.pk).status, 'termine')

    def test_sans_etat_le_programme_ne_change_pas(self):
        self.login(self.coord_pikine)
        self.put(self.url(self.past_activity), {'report': 'Bilan.'})
        self.assertEqual(Program.objects.get(pk=self.pikine_program.pk).status, 'en_cours')

    def test_validations(self):
        self.login(self.coord_pikine)
        self.assertEqual(self.put(self.url(self.past_activity), {}).json()['report'], ["Écrivez le compte rendu de l'activité."])
        self.assertEqual(self.put(self.url(self.past_activity), {'report': '   '}).status_code, 400)
        response = self.put(self.url(self.past_activity), {'report': 'x' * 5001})
        self.assertEqual(response.json()['report'], ['Le compte rendu ne peut pas dépasser 5 000 caractères.'])
        response = self.put(self.url(self.past_activity), {'report': 'Bilan.', 'program_status': 'archive'})
        self.assertEqual(response.json()['program_status'], ['État inconnu. Choisissez Prévu, En cours, Terminé ou Annulé.'])
        self.assertEqual(Activity.objects.get(pk=self.past_activity.pk).report, '')

    def test_le_compte_rendu_se_corrige(self):
        self.login(self.coord_pikine)
        self.put(self.url(self.past_activity), {'report': 'Premier jet.'})
        self.login(self.coord_dakar)
        data = self.put(self.url(self.past_activity), {'report': 'Version corrigée.'}).json()
        self.assertEqual((data['report'], data['report_author']['name']), ('Version corrigée.', 'Mamadou Diop'))

    def test_droits(self):
        self.login(self.coord_thies)
        self.assertEqual(self.put(self.url(self.past_activity), {'report': 'x'}).status_code, 404)
        self.login(self.awa)
        self.assertEqual(self.put(self.url(self.past_activity), {'report': 'x'}).status_code, 403)
        self.client.logout()
        self.assertEqual(self.put(self.url(self.past_activity), {'report': 'x'}).status_code, 401)

    def test_un_membre_ne_lit_pas_le_compte_rendu(self):
        self.login(self.coord_pikine)
        self.put(self.url(self.past_activity), {'report': 'Bilan interne.'})
        self.login(self.awa)
        data = self.client.get(f'/api/member/activities/{self.past_activity.pk}/').json()
        self.assertNotIn('report', data)
        self.assertNotIn('Bilan interne.', str(self.client.get('/api/member/programs/', {'status': 'en_cours'}).json()))


class AttendanceTests(ProgramsTestCase):
    def url(self, activity):
        return f'/api/activities/{activity.pk}/attendance/'

    def setUp(self):
        super().setUp()
        Registration.objects.create(activity=self.past_activity, member=self.awa_member)

    def candidates(self, activity=None):
        self.login(self.coord_pikine)
        data = self.client.get(self.url(activity or self.past_activity)).json()
        return {m['last_name']: m for m in data['members']}

    def test_liste_a_cocher(self):
        members = self.candidates()
        # Les membres actifs de Pikine : Awa (inscrite) et Bineta ; pas Fatou (suspendue), ni Dame (Mbour), ni Cheikh (Rufisque).
        self.assertEqual(sorted(members), ['Fall', 'Ndiaye'])
        self.assertEqual((members['Ndiaye']['registered'], members['Fall']['registered']), (True, False))
        self.assertEqual(members['Ndiaye']['number'], self.awa_member.number)
        self.assertTrue(all(not m['present'] for m in members.values()))

    def test_les_inscrits_d_ailleurs_restent_cochables(self):
        Registration.objects.create(activity=self.past_activity, member=self.cheikh_member)
        self.assertIn('Sarr', self.candidates())

    def test_le_total_des_presents(self):
        self.login(self.coord_pikine)
        response = self.put(self.url(self.past_activity), {'members': [self.awa_member.pk, self.bineta_member.pk]})
        self.assertEqual(response.status_code, 200, response.content)
        data = response.json()
        self.assertEqual((data['noted'], data['present_count']), (True, 2))
        self.assertEqual(data['noted_by']['name'], 'Rokhaya Gueye')
        self.assertTrue(all(m['present'] for m in data['members']))
        activity = self.client.get(f'/api/activities/{self.past_activity.pk}/').json()
        self.assertEqual((activity['attendance_noted'], activity['attendance_count']), (True, 2))
        program = self.client.get(f'/api/programs/{self.pikine_program.pk}/').json()
        counts = {a['id']: a['attendance_count'] for a in program['activities']}
        self.assertEqual(counts[self.past_activity.pk], 2)
        self.assertIsNone(counts[self.future_activity.pk])  # pas encore noté

    def test_l_enregistrement_remplace_la_liste(self):
        self.login(self.coord_pikine)
        self.put(self.url(self.past_activity), {'members': [self.awa_member.pk, self.bineta_member.pk]})
        response = self.put(self.url(self.past_activity), {'members': [self.bineta_member.pk]})
        self.assertEqual(response.json()['present_count'], 1)
        self.assertEqual(list(Attendance.objects.values_list('member_id', flat=True)), [self.bineta_member.pk])

    def test_une_liste_vide_est_une_liste_notee(self):
        self.login(self.coord_pikine)
        data = self.put(self.url(self.past_activity), {'members': []}).json()
        self.assertEqual((data['noted'], data['present_count']), (True, 0))
        self.assertEqual(self.client.get(f'/api/activities/{self.past_activity.pk}/').json()['attendance_count'], 0)

    def test_pas_de_presence_avant_l_activite(self):
        self.login(self.coord_pikine)
        response = self.put(self.url(self.future_activity), {'members': []})
        self.assertEqual((response.status_code, response.json()['code']), (409, 'not_started'))
        self.assertFalse(Attendance.objects.exists())

    def test_on_ne_coche_que_des_membres_du_territoire_ou_inscrits(self):
        self.login(self.coord_pikine)
        for outsider in (self.dame_member, self.fatou_member, self.founder_member):
            response = self.put(self.url(self.past_activity), {'members': [self.awa_member.pk, outsider.pk]})
            self.assertEqual(response.status_code, 400, outsider)
        self.assertFalse(Attendance.objects.exists())
        self.assertEqual(self.put(self.url(self.past_activity), {'members': 'tous'}).status_code, 400)
        self.assertEqual(self.put(self.url(self.past_activity), {'members': [0]}).status_code, 400)
        self.assertEqual(self.put(self.url(self.past_activity), {}).status_code, 400)

    def test_droits(self):
        self.login(self.coord_thies)
        self.assertEqual(self.client.get(self.url(self.past_activity)).status_code, 404)
        self.assertEqual(self.put(self.url(self.past_activity), {'members': []}).status_code, 404)
        self.login(self.awa)
        self.assertEqual(self.client.get(self.url(self.past_activity)).status_code, 403)
        self.assertEqual(self.put(self.url(self.past_activity), {'members': []}).status_code, 403)

    def test_les_presences_sont_notees_par_activite(self):
        self.login(self.coord_pikine)
        self.put(self.url(self.past_activity), {'members': [self.awa_member.pk]})
        other = Activity.objects.create(
            program=self.pikine_program, title='Autre', starts_at=days(-1), place='Mairie', description='Autre activité.'
        )
        self.assertTrue(all(not m['present'] for m in self.client.get(self.url(other)).json()['members']))


class MemberProgramsTests(ProgramsTestCase):
    url = '/api/member/programs/'

    def names(self, **params):
        response = self.client.get(self.url, params)
        self.assertEqual(response.status_code, 200, response.content)
        return [row['name'] for row in response.json()['results']]

    def test_un_membre_voit_les_programmes_de_son_territoire_et_pas_ceux_des_autres(self):
        self.login(self.awa)  # Pikine
        self.assertEqual(self.names(), ['Brigades citoyennes', 'Santé communautaire'])  # sa commune et sa région, par date de début
        self.login(self.dame)  # Mbour
        self.assertEqual(self.names(), ['Éducation civique'])
        self.login(self.cheikh)  # Rufisque : le programme de Pikine ne le concerne pas
        self.assertEqual(self.names(), ['Santé communautaire'])

    def test_par_defaut_les_programmes_prevus_et_en_cours(self):
        Program.objects.filter(pk=self.region_program.pk).update(status=Program.Status.DONE)
        self.login(self.awa)
        self.assertEqual(self.names(), ['Brigades citoyennes'])
        self.assertEqual(self.names(status='termine'), ['Santé communautaire'])
        self.assertEqual(self.names(status='en_cours'), ['Brigades citoyennes'])
        self.assertEqual(self.names(status='prevu'), [])

    def test_les_programmes_annules_sont_caches(self):
        Program.objects.filter(pk=self.pikine_program.pk).update(status=Program.Status.CANCELLED)
        self.login(self.awa)
        for status in ('', 'prevu', 'en_cours', 'termine'):
            self.assertNotIn('Brigades citoyennes', self.names(status=status) if status else self.names())
        self.assertEqual(self.client.get(self.url, {'status': 'annule'}).status_code, 400)
        self.assertEqual(self.client.get(f'/api/member/activities/{self.past_activity.pk}/').status_code, 404)

    def test_recherche(self):
        self.login(self.awa)
        self.assertEqual(self.names(search='brigades'), ['Brigades citoyennes'])
        self.assertEqual(self.names(search='introuvable'), [])

    def test_les_activites_par_ordre_de_date_avec_l_inscription(self):
        Registration.objects.create(activity=self.future_activity, member=self.awa_member)
        Registration.objects.create(activity=self.future_activity, member=self.bineta_member)
        self.login(self.awa)
        program = next(p for p in self.client.get(self.url).json()['results'] if p['name'] == 'Brigades citoyennes')
        self.assertEqual([a['title'] for a in program['activities']], ['Nettoyage de la plage', 'Atelier de quartier'])
        past, future = program['activities']
        self.assertEqual((future['registered'], future['registrations_count'], future['registration_open']), (True, 2, True))
        self.assertEqual((past['registered'], past['registration_open']), (False, False))  # déjà passée
        self.assertEqual(program['manager']['name'], 'Rokhaya Gueye')
        self.login(self.bineta)
        program = next(p for p in self.client.get(self.url).json()['results'] if p['name'] == 'Brigades citoyennes')
        self.assertTrue(program['activities'][1]['registered'])
        self.login(self.cheikh)  # Rufisque : pas ce programme
        self.assertNotIn('Brigades citoyennes', self.names())

    def test_un_fondateur_voit_tout(self):
        self.login(self.founder)
        self.assertEqual(len(self.names()), 3)

    def test_unknown_status_et_acces(self):
        self.login(self.awa)
        self.assertEqual(self.client.get(self.url, {'status': 'archive'}).status_code, 400)
        self.client.logout()
        self.assertEqual(self.client.get(self.url).status_code, 401)

    def test_un_coordonnateur_voit_aussi_les_programmes_de_son_territoire_cote_membre(self):
        self.login(self.coord_pikine)
        self.assertEqual(self.names(), ['Brigades citoyennes', 'Santé communautaire'])

    def test_pages(self):
        self.login(self.founder)
        data = self.client.get(self.url, {'page_size': 2}).json()
        self.assertEqual((data['count'], len(data['results'])), (3, 2))


class RegistrationTests(ProgramsTestCase):
    def url(self, activity):
        return f'/api/member/activities/{activity.pk}/registration/'

    def test_inscription_puis_desinscription(self):
        self.login(self.awa)
        response = self.post(self.url(self.future_activity))
        self.assertEqual(response.status_code, 201, response.content)
        self.assertEqual((response.json()['registered'], response.json()['registrations_count']), (True, 1))
        self.assertTrue(Registration.objects.filter(activity=self.future_activity, member=self.awa_member).exists())
        response = self.client.delete(self.url(self.future_activity))
        self.assertEqual(response.status_code, 200)
        self.assertEqual((response.json()['registered'], response.json()['registrations_count']), (False, 0))
        self.assertFalse(Registration.objects.exists())

    def test_s_inscrire_deux_fois_ne_fait_qu_une_inscription(self):
        self.login(self.awa)
        self.assertEqual(self.post(self.url(self.future_activity)).status_code, 201)
        self.assertEqual(self.post(self.url(self.future_activity)).status_code, 200)
        self.assertEqual(Registration.objects.count(), 1)

    def test_se_desinscrire_sans_etre_inscrit(self):
        self.login(self.awa)
        self.assertEqual(self.client.delete(self.url(self.future_activity)).status_code, 200)

    def test_chacun_s_inscrit_pour_soi(self):
        self.login(self.awa)
        self.post(self.url(self.future_activity))
        self.login(self.bineta)
        self.post(self.url(self.future_activity))
        self.assertEqual(Registration.objects.count(), 2)
        self.client.delete(self.url(self.future_activity))
        self.assertEqual(list(Registration.objects.values_list('member_id', flat=True)), [self.awa_member.pk])

    def test_inscriptions_closes_une_fois_l_activite_commencee(self):
        self.login(self.awa)
        response = self.post(self.url(self.past_activity))
        self.assertEqual((response.status_code, response.json()['code']), (409, 'registration_closed'))
        self.assertEqual(response.json()['detail'], "L'activité a déjà commencé : les inscriptions sont closes.")
        Registration.objects.create(activity=self.past_activity, member=self.awa_member)
        response = self.client.delete(self.url(self.past_activity))
        self.assertEqual(response.status_code, 409)
        self.assertTrue(Registration.objects.filter(activity=self.past_activity).exists())

    def test_programme_termine(self):
        Program.objects.filter(pk=self.pikine_program.pk).update(status=Program.Status.DONE)
        self.login(self.awa)
        response = self.post(self.url(self.future_activity))
        self.assertEqual((response.status_code, response.json()['detail']), (409, "Ce programme n'est plus ouvert aux inscriptions."))

    def test_programme_annule_ou_autre_territoire(self):
        self.login(self.awa)
        self.assertEqual(self.post(self.url(self.thies_activity)).status_code, 404)  # Mbour
        self.assertEqual(self.client.delete(self.url(self.thies_activity)).status_code, 404)
        Program.objects.filter(pk=self.pikine_program.pk).update(status=Program.Status.CANCELLED)
        self.assertEqual(self.post(self.url(self.future_activity)).status_code, 404)
        self.assertFalse(Registration.objects.exists())

    def test_le_programme_d_une_region_concerne_ses_communes(self):
        activity = Activity.objects.create(
            program=self.region_program, title='Dépistage', starts_at=days(3), place='Dakar', description='Dépistage gratuit.'
        )
        self.login(self.cheikh)  # Rufisque, dans la région de Dakar
        self.assertEqual(self.post(self.url(activity)).status_code, 201)
        self.login(self.dame)  # Mbour, région de Thiès
        self.assertEqual(self.post(self.url(activity)).status_code, 404)

    def test_compte_sans_dossier_ou_non_actif(self):
        orphan = User.objects.create_user('sans@example.org', PASSWORD, territory=self.pikine)
        self.login(orphan)
        self.assertEqual(self.post(self.url(self.future_activity)).status_code, 403)
        self.assertEqual(self.client.delete(self.url(self.future_activity)).status_code, 404)
        Member.objects.filter(pk=self.bineta_member.pk).update(status=Member.Status.SUSPENDU)
        self.login(self.bineta)
        self.assertEqual(self.post(self.url(self.future_activity)).status_code, 403)

    def test_connexion_obligatoire(self):
        self.assertEqual(self.post(self.url(self.future_activity)).status_code, 401)

    def test_la_fiche_d_une_activite(self):
        Registration.objects.create(activity=self.future_activity, member=self.awa_member)
        self.login(self.awa)
        data = self.client.get(f'/api/member/activities/{self.future_activity.pk}/').json()
        self.assertEqual((data['title'], data['registered'], data['registration_open']), ('Atelier de quartier', True, True))
        self.assertEqual(data['program']['name'], 'Brigades citoyennes')
        self.login(self.dame)
        self.assertEqual(self.client.get(f'/api/member/activities/{self.future_activity.pk}/').status_code, 404)


class RegistrantsTests(ProgramsTestCase):
    def setUp(self):
        super().setUp()
        Registration.objects.create(activity=self.future_activity, member=self.awa_member)
        Registration.objects.create(activity=self.future_activity, member=self.bineta_member)

    def url(self, activity, suffix=''):
        return f'/api/activities/{activity.pk}/registrations/{suffix}'

    def test_le_coordonnateur_voit_la_liste_des_inscrits(self):
        self.login(self.coord_pikine)
        data = self.client.get(self.url(self.future_activity)).json()
        self.assertEqual(data['count'], 2)
        self.assertEqual([(r['last_name'], r['first_name']) for r in data['results']], [('Fall', 'Bineta'), ('Ndiaye', 'Awa')])
        self.assertEqual(data['results'][1]['email'], 'awa@example.org')
        self.assertEqual(data['results'][1]['number'], self.awa_member.number)
        activity = self.client.get(f'/api/activities/{self.future_activity.pk}/').json()
        self.assertEqual(activity['registrations_count'], 2)

    def test_droits(self):
        self.login(self.coord_thies)
        self.assertEqual(self.client.get(self.url(self.future_activity)).status_code, 404)
        self.login(self.awa)
        self.assertEqual(self.client.get(self.url(self.future_activity)).status_code, 403)
        self.assertEqual(self.client.get(self.url(self.future_activity, 'export/')).status_code, 403)

    def test_export_des_inscrits_note_dans_le_journal(self):
        self.login(self.coord_pikine)
        response = self.client.get(self.url(self.future_activity, 'export/'))
        self.assertEqual(response.status_code, 200)
        rows = [[c.value for c in row] for row in load_workbook(BytesIO(response.content)).active.iter_rows()]
        self.assertEqual(rows[0], ["N° d'adhérent", 'Nom', 'Prénom', 'Courriel', 'Localité', 'Inscrit le'])
        self.assertEqual([(r[1], r[2]) for r in rows[1:]], [('Fall', 'Bineta'), ('Ndiaye', 'Awa')])
        entry = AuditEntry.objects.get()
        self.assertEqual((entry.action, entry.subject), ('liste_exportee', 'Inscrits · Atelier de quartier'))
        self.assertEqual(entry.details['rows'], 2)
        self.assertEqual(entry.details['filters'], {'Programme': 'Brigades citoyennes', 'Activité': 'Atelier de quartier'})


class ProgramExportTests(ProgramsTestCase):
    url = '/api/programs/export/'

    def rows(self, **params):
        response = self.client.get(self.url, params)
        self.assertEqual(response.status_code, 200, response.content)
        return [[c.value for c in row] for row in load_workbook(BytesIO(response.content)).active.iter_rows()]

    def test_le_fichier_est_exactement_la_liste_filtree(self):
        self.login(self.founder)
        for params in ({}, {'status': 'en_cours'}, {'territory': self.dakar.pk}, {'search': 'écoles'}, {'ordering': 'name'}):
            with self.subTest(params=params):
                listed = [r['name'] for r in self.client.get('/api/programs/', {**params, 'page_size': 100}).json()['results']]
                self.assertEqual([row[0] for row in self.rows(**params)[1:]], listed)

    def test_colonnes_et_perimetre(self):
        self.login(self.coord_pikine)
        rows = self.rows()
        self.assertEqual(rows[0], ['Programme', 'Objectif', 'Localité', 'Responsable', 'Début', 'Fin', 'État', "Nombre d'activités"])
        self.assertEqual(len(rows), 2)
        self.assertEqual((rows[1][0], rows[1][2], rows[1][3], rows[1][6], rows[1][7]), ('Brigades citoyennes', 'Pikine', 'Rokhaya Gueye', 'En cours', 2))

    def test_note_dans_le_journal(self):
        self.login(self.coord_dakar)
        self.client.get(self.url, {'status': 'en_cours'})
        entry = AuditEntry.objects.get()
        self.assertEqual((entry.subject, entry.details['rows'], entry.details['filters']), ('Programmes', 1, {'État': 'En cours'}))

    def test_reserve_au_bureau(self):
        self.login(self.awa)
        self.assertEqual(self.client.get(self.url).status_code, 403)
        self.assertFalse(AuditEntry.objects.exists())
